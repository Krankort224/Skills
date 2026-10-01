"""Generate one semantic sample in each preset; render separately."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from PIL import Image, ImageDraw
from document_tools import (add_data_table, add_field, add_figure, add_list,
                            add_omml, add_text, apply_preset, load_preset, package_report, resolve_fonts)


def make_chart(path):
    image = Image.new("RGB", (1200, 430), "white"); draw = ImageDraw.Draw(image)
    draw.line([(100, 50), (100, 350), (1130, 350)], fill="#777777", width=3)
    values = [40, 55, 70]
    for i, val in enumerate(values):
        x = 200 + i * 290
        draw.rectangle((x, 350 - val * 4, x + 150, 349), fill="#31566B")
        draw.text((x + 62, 365), str(i + 1), fill="black", font_size=26)
        draw.text((x + 60, 325 - val * 4), str(val), fill="black", font_size=24)
    image.save(path)


def build(name, out, chart):
    config = load_preset(name)
    import shutil
    if shutil.which("fc-list"):
        config, substitutions = resolve_fonts(config)
        if substitutions: print(name + " font substitutions: " + str(substitutions), flush=True)
    doc = Document(); apply_preset(doc, config)
    add_text(doc, "Технический отчёт", "Title")
    add_text(doc, "Образец оформления документа", "Subtitle")
    add_text(doc, "Версия 1   Демонстрационные данные", "Metadata")
    add_text(doc, "Исходные данные", "Heading 1")
    add_text(doc, "Документ показывает работу именованных стилей, редактируемых таблиц и настоящих списков. Числовые значения служат только примером.")
    add_text(doc, "Параметры", "Heading 2")
    width = config["page"]["width_mm"] - config["page"]["margins_mm"]["left"] - config["page"]["margins_mm"]["right"]
    add_data_table(doc, ["Показатель", "Значение", "Единица"], [["Расход", "660", "м³/ч"], ["Температура", "23", "°C"], ["Датчик", "O₂", "—"]], [width * .60, width * .20, width * .20], config)
    add_text(doc, "Порядок проверки", "Heading 3")
    _, nid = add_list(doc, "Проверить исходные данные", ordered=True, config=config)
    add_list(doc, "Сопоставить единицы измерения", level=1, ordered=True, num_id=nid, config=config)
    add_list(doc, "Проверить результат и вёрстку", ordered=True, num_id=nid, config=config)
    add_text(doc, "Допущение", "Heading 4")
    add_text(doc, "Все значения в образце демонстрационные.", "Note")
    heading = add_text(doc, "Результаты проверки", "Heading 1"); heading.paragraph_format.page_break_before = True
    p, _ = add_figure(doc, chart, min(150, width), alt="Столбцы демонстрационных значений 40 55 и 70")
    p.paragraph_format.keep_with_next = True
    add_text(doc, "Рисунок 1  Демонстрационные значения", "Caption")
    add_text(doc, "Расчётная зависимость", "Heading 2")
    p = doc.add_paragraph()
    add_omml(p, '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>Q = vA</m:t></m:r></m:oMath>')
    add_text(doc, "Фрагмент кода", "Heading 3")
    add_text(doc, "flow = velocity * area\nassert flow > 0", "Code Block")
    add_text(doc, "Примечания к проверке", "Heading 5")
    add_text(doc, "Формула остаётся редактируемой в Word.", "Note")
    add_text(doc, "Ссылка на источник", "Heading 6")
    add_text(doc, [{"text": "Документация WordprocessingML", "link": "https://learn.microsoft.com/en-us/office/open-xml/word/structure-of-a-wordprocessingml-document"}])
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Страница "); add_field(footer, "PAGE", "1"); footer.add_run(" из "); add_field(footer, "NUMPAGES", "2")
    for run in footer.runs:
        run.font.size = Pt(config["footer"]["size_pt"]); run.font.color.rgb = RGBColor.from_string(config["footer"]["color"])
    out.mkdir(parents=True, exist_ok=True); path = out / (name + ".docx"); doc.save(path)
    report = package_report(path)
    if report["errors"]: raise ValueError(report["errors"])
    print(str(path), flush=True)


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--out", type=Path, required=True); args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True); chart = args.out / "demo-chart.png"; make_chart(chart)
    for name in ("technical", "corporate", "minimal", "academic"):
        build(name, args.out, chart)


if __name__ == "__main__": main()
