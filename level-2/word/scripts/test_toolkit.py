"""Focused behavior checks for supported components and corruption detection."""
import hashlib
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree
from PIL import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from document_tools import (NS, add_data_table, add_field, add_figure, add_layout_table,
    add_list, add_omml, add_text, apply_preset, body_text, load_preset, package_report, source_coverage, resolve_fonts)


def run():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp); checks = 0
        from preset_markdown import parse_markdown, parameter_table
        from extract_preset import extract_preset
        from docx.shared import Mm, Pt
        expected_hashes = {'minimal': '0af99e06cb0242f48ec7c62f3badcb82bf87d7d9dfe3dd226478ee119cc5e9d0', 'technical': 'eaaf45e3ab481bc3040778407d47dbf02dfa25cf9f2c5a69e663389730489b6d', 'corporate': '5efc63ec627775b71f26c18ac0d1da779be27a70123c2b52875b8629a93e95f2', 'academic': '299e6b5d5e8155f6aeb2f8ccb9d293bacb2b8f331f0c12b63b7a895d6b0ba862'}
        for name, digest in expected_hashes.items():
            config = load_preset(name)
            assert hashlib.sha256(json.dumps(config, sort_keys=True, separators=(',', ':')).encode()).hexdigest() == digest
            assert parse_markdown(parameter_table(config)) == config
            md = root / (name + '.md'); md.write_text('Status: ready\n\n' + parameter_table(config))
            assert load_preset(md) == config
            applied = Document(); apply_preset(applied, config)
            saved = root / ('geometry-' + name + '.docx'); applied.save(saved); applied = Document(saved)
            assert abs(applied.sections[0].page_width.mm - config['page']['width_mm']) < .02
            for side, expected in config['page']['margins_mm'].items():
                assert abs(getattr(applied.sections[0], side + '_margin').mm - expected) < .02
            for role, spec in config['styles'].items():
                assert applied.styles[role].font.name == spec['font']
                assert applied.styles[role].font.size.pt == spec['size_pt']
                assert str(applied.styles[role].font.color.rgb) == spec['color']
            checks += 5
        valid = parameter_table(load_preset('technical'))
        bad_tables = [valid + '| name | x | string | - |\n',
            valid.replace('| page.width_mm | 210 | number | mm |', '| page.width_mm | 210 | number | pt |'),
            valid.replace('| version | 1 | number | - |', '| version | 1 | integer | - |'),
            valid.replace('| page.width_mm | 210 | number | mm |', '| page.width_mm | 1e999 | number | mm |'),
            valid.replace('| page.width_mm | 210 | number | mm |', '| page.width_mm | -1 | number | mm |'),
            valid.replace('| name | technical | string | - |', ''),
            valid.replace('| styles.Normal.size_pt | 11 | number | pt |', '| styles.Normal.size_pt | 11 | string | pt |'),
            valid + '| page | conflict | string | - |\n',
            valid.replace('| table.repeat_header | true | boolean | - |', '| table.repeat_header | yes | boolean | - |')]
        for invalid in bad_tables:
            try: parse_markdown(invalid)
            except ValueError: checks += 1
            else: raise AssertionError('Invalid preset accepted')
        escaped = load_preset('technical'); escaped['name'] = 'A|B\\C'
        assert parse_markdown(parameter_table(escaped)) == escaped; checks += 1
        source = Document(); source.sections[0].left_margin = Mm(31)
        source.styles['Heading 1'].font.size = Pt(19)
        paragraph = source.add_paragraph('PRIVATE SOURCE TEXT', 'Heading 1'); paragraph.runs[0].bold = True
        source.styles.add_style('Corp.Body', 1)
        reference = root / 'reference.docx'; source.save(reference)
        draft, report = extract_preset(reference, 'technical', 0)
        assert 'PRIVATE SOURCE TEXT' not in draft
        assert 'page.margins_mm.left' in report['observed_paths']
        assert 'palette.accent' in report['defaulted_paths']
        assert report['direct_formatting'][0]['run_properties'] == ['b']
        assert 'Corp.Body' in report['omitted_style_names']
        draft_path = root / 'draft.md'; draft_path.write_text(draft)
        borrowed = load_preset(draft_path, allow_draft=True)
        assert abs(borrowed['page']['margins_mm']['left'] - 31) < .02
        assert borrowed['styles']['Heading 1']['size_pt'] == 19
        assert extract_preset(reference, 'technical', 0) == (draft, report)
        try: load_preset(draft_path)
        except ValueError as exc: assert 'render/compare' in str(exc)
        else: raise AssertionError('Draft accepted')
        try: extract_preset(reference, 'technical', 1)
        except ValueError: pass
        else: raise AssertionError('Missing section accepted')
        cli_path = root / 'cli-draft.md'
        cli = subprocess.run([sys.executable, str(Path(__file__).with_name('document_tool.py')), 'extract-preset', str(reference), '--base', 'technical', '--out', str(cli_path)], capture_output=True, text=True, check=True)
        assert cli_path.read_text() == draft
        assert json.loads(cli.stdout) == report
        checks += 11
        original = load_preset("technical")
        effective, substitutions = resolve_fonts(original, ["Nimbus Sans", "DejaVu Sans Mono"])
        assert effective["styles"]["Code Block"]["font"] == "DejaVu Sans Mono"
        assert original["styles"]["Code Block"]["font"] == "Consolas"; checks += 2
        for name in ("technical", "corporate", "minimal", "academic"):
            config = load_preset(name); doc = Document(); apply_preset(doc, config)
            add_text(doc, "Проверка кириллицы O₂ м³", "Heading 1")
            _, nid = add_list(doc, "Первый пункт", ordered=True, config=config)
            add_list(doc, "Вложенный пункт", level=1, ordered=True, num_id=nid, config=config)
            grid = add_layout_table(doc, [60, 90]); add_list(grid.cell(0, 0), "Пункт в ячейке")
            add_data_table(doc, ["Имя", "Значение"], [["Строка", "42"]], [100, 50], config)
            image = root / "figure.png"; Image.new("RGB", (200, 100), "white").save(image)
            _, shape = add_figure(doc, image, 80, max_height_mm=20)
            assert abs(shape.width.mm - 40) < .02 and abs(shape.height.mm - 20) < .02
            p = doc.add_paragraph(); add_omml(p, '<m:oMath xmlns:m="' + NS["m"] + '"><m:r><m:t>x = 1</m:t></m:r></m:oMath>')
            add_field(doc.sections[0].footer.paragraphs[0], "PAGE", "1")
            simple = OxmlElement("w:fldSimple"); simple.set(qn("w:instr"), "PAGE")
            doc.sections[0].header.paragraphs[0]._p.append(simple)
            path = root / (name + ".docx"); doc.save(path)
            report = package_report(path); assert not report["errors"], report["errors"]
            assert report["counts"]["math"] == 1 and report["counts"]["drawings"] == 1
            assert report["counts"]["fields"] == 2
            assert report["story_counts"]["word/footer1.xml"]["fields"] == 1; checks += 2
            assert not source_coverage(path, ["Первый пункт", "Вложенный пункт", "Пункт в ячейке"])["missing"]
            assert source_coverage(path, ["Отсутствующий блок"])["missing"]
            assert source_coverage(path, ["Первый пункт", "Первый пункт"])["missing"]
            checks += 7
        before = body_text(path); apply_preset(doc, load_preset("minimal")); styled = root / "styled.docx"; doc.save(styled)
        assert before == body_text(styled); checks += 1
        try: add_field(doc.add_paragraph(), 'DDE external target')
        except ValueError: checks += 1
        else: raise AssertionError("External field not rejected")
        broken = root / "broken.docx"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(broken, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    xml = etree.fromstring(data)
                    for rel in xml:
                        if rel.get("Type", "").endswith("/image"): rel.set("Target", "media/missing.png")
                    data = etree.tostring(xml)
                target.writestr(item, data)
        assert any("missing target" in error for error in package_report(broken)["errors"]); checks += 1
        p, _ = add_list(doc.sections[0].footer, "Footer list")
        p._p.xpath(".//w:numId")[0].set(qn("w:val"), "99999")
        invalid_footer = root / "invalid-footer.docx"; doc.save(invalid_footer)
        assert any("word/footer1.xml missing numbering ID" in error for error in package_report(invalid_footer)["errors"]); checks += 1
        print(f"Passed {checks} behavior checks across four presets")


if __name__ == "__main__": run()

