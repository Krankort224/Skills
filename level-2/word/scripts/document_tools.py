"""Reusable editable DOCX components and conservative package checks."""
from __future__ import annotations

import collections
import copy
import io
import json
import posixpath
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


def load_preset(name_or_path, allow_draft=False):
    """Load one self-contained Markdown preset by built-in name or .md path."""
    from preset_markdown import parse_markdown
    path = Path(name_or_path)
    if not path.is_file():
        path = ROOT / "assets" / "presets" / (str(name_or_path) + ".md")
    if path.suffix.lower() != ".md":
        raise ValueError("Presets must be single Markdown files")
    text = path.read_text(encoding="utf-8")
    statuses = re.findall(r"^Status: (ready|draft)$", text, re.MULTILINE)
    if len(statuses) != 1: raise ValueError("Preset requires exactly one Status: ready or Status: draft")
    if statuses[0] == "draft" and not allow_draft:
        raise ValueError("Draft preset: review unobserved defaults and render/compare the result before changing Status: draft to Status: ready")
    return parse_markdown(text)


def apply_preset(doc, config, geometry=True):
    if geometry:
        for section in doc.sections:
            section.page_width = Mm(config["page"]["width_mm"])
            section.page_height = Mm(config["page"]["height_mm"])
            for side, value in config["page"]["margins_mm"].items():
                setattr(section, side + "_margin", Mm(value))
            section.header_distance = Mm(config["page"]["header_mm"])
            section.footer_distance = Mm(config["page"]["footer_mm"])
    for name, spec in config["styles"].items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = spec["font"]
        style.font.size = Pt(spec["size_pt"])
        style.font.bold = spec.get("bold", False)
        style.font.italic = spec.get("italic", False)
        style.font.color.rgb = RGBColor.from_string(spec["color"])
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(qn("w:" + attr), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn("w:" + attr), spec["font"])
        pf = style.paragraph_format
        pf.space_before = Pt(spec.get("before_pt", 0))
        pf.space_after = Pt(spec.get("after_pt", 0))
        pf.line_spacing = spec.get("line_spacing", config["styles"]["Normal"]["line_spacing"])
        pf.keep_with_next = spec.get("keep_with_next", False)
        pf.widow_control = True
        props = style.element.get_or_add_pPr()
        border = props.find(qn("w:pBdr"))
        if border is not None:
            props.remove(border)
    return doc


def resolve_fonts(config, available_fonts=None):
    """Resolve a task-local preset copy against an explicit/Fontconfig inventory."""
    if available_fonts is None:
        if not shutil.which("fc-list"):
            raise RuntimeError("Provide available_fonts explicitly when Fontconfig is unavailable")
        listed = subprocess.run(["fc-list", "-f", "%{family}\n"], check=True, capture_output=True, text=True).stdout
        available_fonts = [name.strip() for line in listed.splitlines() for name in line.split(",")]
    inventory = {name.casefold(): name for name in available_fonts}
    candidates = {"Arial": ["Liberation Sans", "Nimbus Sans", "DejaVu Sans"],
                  "Calibri": ["Carlito", "DejaVu Sans", "Nimbus Sans"],
                  "Times New Roman": ["Liberation Serif", "Nimbus Roman", "DejaVu Serif"],
                  "Consolas": ["Liberation Mono", "DejaVu Sans Mono", "Courier New"]}
    effective, substitutions = copy.deepcopy(config), {}
    for spec in effective["styles"].values():
        requested = spec["font"]
        choices = [requested, config["fonts"]["fallbacks"].get(requested, requested)] + candidates.get(requested, [])
        matched = next((inventory[n.casefold()] for n in choices if n.casefold() in inventory), None)
        if matched is None:
            raise ValueError("No available font for " + requested)
        spec["font"] = matched
        if matched.casefold() != requested.casefold(): substitutions[requested] = matched
    return effective, substitutions


def add_text(parent, text_or_spans, style="Normal"):
    paragraph = parent.add_paragraph(style=style)
    spans = [{"text": text_or_spans}] if isinstance(text_or_spans, str) else text_or_spans
    for span in spans:
        run = paragraph.add_run(span["text"])
        if span.get("bold"):
            run.bold = True
        if span.get("italic"):
            run.italic = True
        if span.get("code"):
            run.font.name = "Consolas"
        if span.get("link"):
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            rid = paragraph.part.relate_to(span["link"], RT.HYPERLINK, is_external=True)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), rid)
            run._r.getparent().remove(run._r)
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
    return paragraph


def _next_id(root, tag, attr):
    ids = [int(e.get(qn(attr))) for e in root.findall(qn(tag))]
    return max(ids, default=-1) + 1


def create_numbering(doc, ordered=False, config=None):
    root = doc.part.numbering_part.element
    aid = _next_id(root, "w:abstractNum", "w:abstractNumId")
    nid = _next_id(root, "w:num", "w:numId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    for level in range(9):
        lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), str(level))
        for tag, val in (("start", "1"), ("numFmt", "decimal" if ordered else "bullet"),
                         ("lvlText", "%" + str(level + 1) + "." if ordered else "•"), ("lvlJc", "left"), ("suff", "tab")):
            node = OxmlElement("w:" + tag); node.set(qn("w:val"), val); lvl.append(node)
        settings = config["list"] if config else {"indent_mm": 5, "hanging_mm": 3, "level_step_mm": 6}
        pp = OxmlElement("w:pPr"); ind = OxmlElement("w:ind")
        left = str(round((settings["indent_mm"] + level * settings["level_step_mm"]) * 1440 / 25.4))
        ind.set(qn("w:left"), left)
        ind.set(qn("w:hanging"), str(round(settings["hanging_mm"] * 1440 / 25.4)))
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), left); tabs.append(tab)
        pp.append(tabs); pp.append(ind); lvl.append(pp); abstract.append(lvl)
    root.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(aid)); num.append(ref); root.append(num)
    return nid


def add_list(parent, text, level=0, ordered=False, num_id=None, config=None, style="Normal"):
    if level not in range(9):
        raise ValueError("List level must be 0–8")
    if hasattr(parent, "sections"):
        doc = parent
    else:
        class Owner:
            part = parent.part.package.main_document_part
        doc = Owner()
    num_id = create_numbering(doc, ordered, config) if num_id is None else num_id
    p = add_text(parent, text, style)
    num = p._p.get_or_add_pPr().get_or_add_numPr()
    num.get_or_add_ilvl().val = level; num.get_or_add_numId().val = num_id
    if config:
        p.paragraph_format.space_after = Pt(config["list"]["after_pt"])
    return p, num_id


def set_cell_padding(cell, **values_mm):
    props = cell._tc.get_or_add_tcPr()
    old = props.find(qn("w:tcMar"))
    if old is not None:
        props.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, value in values_mm.items():
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), str(round(value * 1440 / 25.4))); node.set(qn("w:type"), "dxa"); mar.append(node)
    props.append(mar)


def set_table_borders(table, color="D9D9D9", width_pt=0.5, visible=True):
    props = table._tbl.tblPr
    old = props.find(qn("w:tblBorders"))
    if old is not None:
        props.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side); e.set(qn("w:val"), "single" if visible else "nil")
        e.set(qn("w:sz"), str(round(width_pt * 8))); e.set(qn("w:color"), color); borders.append(e)
    props.append(borders)


def _set_widths(table, widths):
    if not widths or any(w <= 0 for w in widths):
        raise ValueError("Table widths must be positive")
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW")); width.set(qn("w:w"), str(round(sum(widths) * 1440 / 25.4))); width.set(qn("w:type"), "dxa")


def add_data_table(doc, headers, rows, widths_mm, config):
    if len(headers) != len(widths_mm) or any(len(row) != len(headers) for row in rows):
        raise ValueError("Inconsistent table columns")
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers)); _set_widths(table, widths_mm)
    set_table_borders(table, config["palette"]["border"], config["table"]["border_pt"])
    for index, values in enumerate([headers] + list(rows)):
        row = table.rows[index]
        if index == 0 and config["table"]["repeat_header"]:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, value in zip(row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_padding(cell, **config["table"]["padding_mm"])
            p = cell.paragraphs[0]; p.style = "Table Text"; run = p.add_run(str(value))
            if index == 0:
                run.bold = True; run.font.color.rgb = RGBColor.from_string(config["palette"]["table_header_text"])
            fill = config["palette"]["table_header"] if index == 0 else (config["palette"]["alternating_row"] if index % 2 == 0 else "FFFFFF")
            shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(shading)
    return table


def add_layout_table(doc, widths_mm, gutter_mm=6):
    table = doc.add_table(rows=1, cols=len(widths_mm)); _set_widths(table, widths_mm)
    set_table_borders(table, visible=False)
    for i, cell in enumerate(table.rows[0].cells):
        set_cell_padding(cell, top=0, bottom=0, left=0 if i == 0 else gutter_mm / 2,
                         right=0 if i == len(widths_mm) - 1 else gutter_mm / 2)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return table


def add_figure(parent, path, width_mm, max_height_mm=None, alt=""):
    if width_mm <= 0:
        raise ValueError("Figure width must be positive")
    with Image.open(path) as img:
        width_px, height_px = img.size
    height_mm = width_mm * height_px / width_px
    if max_height_mm is not None and height_mm > max_height_mm:
        if max_height_mm <= 0:
            raise ValueError("Figure height limit must be positive")
        width_mm *= max_height_mm / height_mm
    p = parent.add_paragraph()
    shape = p.add_run().add_picture(str(path), width=Mm(width_mm))
    shape._inline.docPr.set("descr", alt)
    return p, shape


def add_field(paragraph, instruction, cached=""):
    if not re.match(r"^(PAGE|NUMPAGES|TOC|SEQ|REF|PAGEREF)(?:\s|$)", instruction) or re.search(r"[\r\n]", instruction):
        raise ValueError("Only local Word field instructions are supported")
    for tag, typ in (("fldChar", "begin"), ("instrText", None), ("fldChar", "separate")):
        run = OxmlElement("w:r"); node = OxmlElement("w:" + tag)
        if typ:
            node.set(qn("w:fldCharType"), typ)
            if typ == "begin": node.set(qn("w:dirty"), "true")
        else:
            node.set(qn("xml:space"), "preserve"); node.text = instruction
        run.append(node); paragraph._p.append(run)
    paragraph.add_run(cached)
    run = OxmlElement("w:r"); end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); run.append(end); paragraph._p.append(run)


def add_omml(paragraph, xml):
    root = etree.fromstring(xml.encode() if isinstance(xml, str) else xml, PARSER)
    if root.tag not in (qn("m:oMath"), qn("m:oMathPara")):
        raise ValueError("Expected native oMath or oMathPara root")
    paragraph._p.append(root)


def _texts(root):
    return ["".join(p.xpath(".//w:t/text() | .//m:t/text()", namespaces=NS)) for p in root.xpath(".//w:p", namespaces=NS)]


def body_text(path):
    with zipfile.ZipFile(path) as package:
        return _texts(etree.fromstring(package.read("word/document.xml"), PARSER))


def normalize_text(text):
    return " ".join(text.replace("\u00a0", " ").split())


def source_coverage(path, expected_lines):
    actual = collections.Counter(normalize_text(line) for line in body_text(path) if normalize_text(line))
    expected = collections.Counter(normalize_text(line) for line in expected_lines if normalize_text(line))
    missing = {text: count - actual[text] for text, count in expected.items() if actual[text] < count}
    extra = {text: actual[text] - count for text, count in expected.items() if actual[text] > count}
    return {"expected_blocks": sum(expected.values()), "missing": missing, "extra_occurrences": extra}


def package_report(path):
    errors, warnings, xmls, rels = [], [], {}, {}
    with zipfile.ZipFile(path) as package:
        names = package.namelist(); files = set(names)
        if len(files) != len(names): errors.append("Duplicate ZIP part names")
        corrupt = package.testzip()
        if corrupt: errors.append("Corrupt ZIP entry: " + corrupt)
        for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml"):
            if required not in files: errors.append("Missing required part: " + required)
        for name in names:
            if name.endswith((".xml", ".rels")):
                try: xmls[name] = etree.fromstring(package.read(name), PARSER)
                except etree.XMLSyntaxError as exc: errors.append(name + ": " + str(exc))
        for name, root in xmls.items():
            if not name.endswith(".rels"): continue
            owner = "" if name == "_rels/.rels" else posixpath.join(posixpath.dirname(posixpath.dirname(name)), posixpath.basename(name)[:-5])
            ownerrels = {}
            for rel in root:
                rid = rel.get("Id"); target = rel.get("Target", ""); ownerrels[rid] = target
                if rel.get("TargetMode") == "External": continue
                dest = unquote(urlparse(target).path)
                resolved = posixpath.normpath(dest.lstrip("/")) if dest.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(owner), dest))
                if resolved not in files: errors.append(name + " missing target: " + resolved)
            rels[owner] = ownerrels
        for name, root in xmls.items():
            if name.endswith(".rels"): continue
            for element in root.iter():
                for attr, value in element.attrib.items():
                    if attr.startswith("{" + NS["r"] + "}") and value not in rels.get(name, {}):
                        errors.append(name + " unresolved relationship: " + value)
        types = xmls.get("[Content_Types].xml")
        if types is not None:
            defaults = {e.get("Extension") for e in types if e.tag.endswith("Default")}
            overrides = {e.get("PartName", "").lstrip("/") for e in types if e.tag.endswith("Override")}
            for name in files:
                if name.endswith("/") or name == "[Content_Types].xml": continue
                if name not in overrides and name.rsplit(".", 1)[-1] not in defaults: errors.append("Missing content type: " + name)
        body = xmls.get("word/document.xml")
        story_roots = {name: root for name, root in xmls.items()
                       if re.match(r"word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$", name)}
        structure_paths = {"paragraphs":".//w:p", "tables":".//w:tbl", "drawings":".//w:drawing",
                           "math":".//m:oMath", "fields":".//w:fldSimple | .//w:fldChar[@w:fldCharType='begin']",
                           "text_boxes":".//w:txbxContent", "revisions":".//w:ins | .//w:del",
                           "content_controls":".//w:sdt", "sections":".//w:sectPr"}
        story_counts = {name: {label: len(root.xpath(xpath, namespaces=NS))
                              for label, xpath in structure_paths.items()} for name, root in story_roots.items()}
        counts = {}
        if body is not None:
            for label in structure_paths:
                counts[label] = sum(values[label] for values in story_counts.values())
            counts["media"] = sum(name.startswith("word/media/") for name in files)
            numroot = xmls.get("word/numbering.xml")
            nums = set(numroot.xpath("./w:num/@w:numId", namespaces=NS)) if numroot is not None else set()
            abstracts = set(numroot.xpath("./w:abstractNum/@w:abstractNumId", namespaces=NS)) if numroot is not None else set()
            for name, root in story_roots.items():
                for val in root.xpath(".//w:numPr/w:numId/@w:val", namespaces=NS):
                    if val != "0" and val not in nums: errors.append(name + " missing numbering ID: " + val)
            if numroot is not None:
                for val in numroot.xpath("./w:num/w:abstractNumId/@w:val", namespaces=NS):
                    if val not in abstracts: errors.append("Missing abstract numbering ID: " + val)
            if counts["revisions"] or counts["text_boxes"] or counts["content_controls"]:
                warnings.append("Complex objects present; inspect before mutation")
        stories = {name: _texts(root) for name, root in story_roots.items()}
        geometry, styles = [], []
        try:
            doc = Document(path)
            geometry = [{"width_mm": round(s.page_width.mm, 2), "height_mm": round(s.page_height.mm, 2),
                         "margins_mm": {k: round(getattr(s, k + "_margin").mm, 2) for k in ("left","right","top","bottom")}} for s in doc.sections]
            styles = [s.name for s in doc.styles]
        except Exception as exc:
            errors.append("High-level reopen failed: " + str(exc))
    return {"errors": sorted(set(errors)), "warnings": warnings, "counts": counts,
            "geometry": geometry, "stories": stories, "story_counts": story_counts, "styles": styles}

