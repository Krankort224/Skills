"""Focused behavior checks for supported components and corruption detection."""
import tempfile
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree
from PIL import Image
from document_tools import (NS, add_data_table, add_field, add_figure, add_layout_table,
    add_list, add_omml, add_text, apply_preset, body_text, load_preset, package_report, source_coverage, resolve_fonts)


def run():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp); checks = 0
        original = load_preset("technical")
        effective, substitutions = resolve_fonts(original, ["Nimbus Sans", "DejaVu Sans Mono"])
        assert effective["styles"]["Code Block"]["font"] == "DejaVu Sans Mono"
        assert original["styles"]["Code Block"]["font"] == "Consolas"; checks += 2
        for name in ("technical", "corporate", "minimal", "academic"):
            config = load_preset(name); doc = Document(); apply_preset(doc, config)
            add_text(doc, "Проверка кириллицы O₂ м³", "Heading 1")
            _, nid = add_list(doc, "Первый пункт", ordered=True, config=config)
            add_list(doc, "Вложенный пункт", level=1, ordered=True, num_id=nid, config=config)
            grid = add_layout_table(doc, [60, 90]); add_list(grid.cell(0, 0), "Пункт в ячейке")
            add_data_table(doc, ["Имя", "Значение"], [["Строка", "42"]], [100, 50], config)
            image = root / "figure.png"; Image.new("RGB", (200, 100), "white").save(image)
            _, shape = add_figure(doc, image, 80, max_height_mm=20)
            assert abs(shape.width.mm - 40) < .02 and abs(shape.height.mm - 20) < .02
            p = doc.add_paragraph(); add_omml(p, '<m:oMath xmlns:m="' + NS["m"] + '"><m:r><m:t>x = 1</m:t></m:r></m:oMath>')
            add_field(doc.sections[0].footer.paragraphs[0], "PAGE", "1")
            path = root / (name + ".docx"); doc.save(path)
            report = package_report(path); assert not report["errors"], report["errors"]
            assert report["counts"]["math"] == 1 and report["counts"]["drawings"] == 1
            assert not source_coverage(path, ["Первый пункт", "Вложенный пункт", "Пункт в ячейке"])["missing"]
            assert source_coverage(path, ["Отсутствующий блок"])["missing"]
            assert source_coverage(path, ["Первый пункт", "Первый пункт"])["missing"]
            checks += 7
        before = body_text(path); apply_preset(doc, load_preset("minimal")); styled = root / "styled.docx"; doc.save(styled)
        assert before == body_text(styled); checks += 1
        try: add_field(doc.add_paragraph(), 'DDE external target')
        except ValueError: checks += 1
        else: raise AssertionError("External field not rejected")
        broken = root / "broken.docx"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(broken, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    xml = etree.fromstring(data)
                    for rel in xml:
                        if rel.get("Type", "").endswith("/image"): rel.set("Target", "media/missing.png")
                    data = etree.tostring(xml)
                target.writestr(item, data)
        assert any("missing target" in error for error in package_report(broken)["errors"]); checks += 1
        print(f"Passed {checks} behavior checks across four presets")


if __name__ == "__main__": run()
